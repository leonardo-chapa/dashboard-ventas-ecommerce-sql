"""Construye el dashboard en Excel con fórmulas y gráficas nativas."""

from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AZUL, AZUL_OSC, GRIS, TINTA2 = "2A78D6", "1C5CAB", "F2F2EF", "52514E"
F = "Arial"
ESTADOS = {
    "AC": "Acre", "AL": "Alagoas", "AM": "Amazonas", "AP": "Amapá", "BA": "Bahía", "CE": "Ceará",
    "DF": "Distrito Federal", "ES": "Espírito Santo", "GO": "Goiás", "MA": "Maranhão",
    "MG": "Minas Gerais", "MS": "Mato Grosso do Sul", "MT": "Mato Grosso", "PA": "Pará",
    "PB": "Paraíba", "PE": "Pernambuco", "PI": "Piauí", "PR": "Paraná", "RJ": "Río de Janeiro",
    "RN": "Río Grande del Norte", "RO": "Rondônia", "RR": "Roraima", "RS": "Río Grande del Sur",
    "SC": "Santa Catarina", "SE": "Sergipe", "SP": "São Paulo", "TO": "Tocantins",
}
FUENTE = ("Fuente: Olist, Brazilian E-Commerce Public Dataset (Kaggle, CC BY-NC-SA 4.0). "
          "Pedidos entregados, ene-2017 a ago-2018. Montos en reales brasileños (BRL).")


def _encabezado(ws, fila, columnas):
    for j, texto in enumerate(columnas, start=1):
        c = ws.cell(fila, j, texto)
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color=AZUL_OSC)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[fila].height = 32


def _tabla(ws, df, formatos, anchos, titulo, nota=None):
    ws["A1"] = titulo
    ws["A1"].font = Font(name=F, bold=True, size=13)
    if nota:
        ws["A2"] = nota
        ws["A2"].font = Font(name=F, italic=True, size=9, color=TINTA2)
    _encabezado(ws, 4, list(df.columns))
    for i, fila in enumerate(df.itertuples(index=False), start=5):
        for j, v in enumerate(fila, start=1):
            c = ws.cell(i, j, None if pd.isna(v) else (v.item() if hasattr(v, "item") else v))
            c.font = Font(name=F)
            if formatos.get(j):
                c.number_format = formatos[j]
    for j, a in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(j)].width = a
    ws.freeze_panes = "A5"
    return 5, 4 + len(df)   # primera y última fila de datos


def construir(r: dict[str, pd.DataFrame], ruta: str):
    wb = Workbook()
    dash = wb.active
    dash.title = "Dashboard"

    # ---------------- Ventas mensuales (crecimiento con fórmula)
    ws = wb.create_sheet("Ventas mensuales")
    vm = r["ventas_mensuales"][["mes", "pedidos", "ingreso_brl"]].rename(
        columns={"mes": "Mes", "pedidos": "Pedidos", "ingreso_brl": "Ingreso (BRL)"})
    i0, i1 = _tabla(ws, vm, {2: "#,##0", 3: "#,##0"}, [12, 12, 16, 16, 18],
                    "Ventas mensuales", "Ticket y crecimiento calculados con fórmulas de Excel.")
    for c, t in [(4, "Ticket promedio (BRL)"), (5, "Crecimiento vs. mes anterior")]:
        ws.cell(4, c, t)
    _encabezado(ws, 4, ["Mes", "Pedidos", "Ingreso (BRL)", "Ticket promedio (BRL)", "Crecimiento vs. mes anterior"])
    for i in range(i0, i1 + 1):
        ws.cell(i, 4, f"=IFERROR(C{i}/B{i},0)").number_format = "#,##0.00"
        ws.cell(i, 5, "" if i == i0 else f"=IFERROR(C{i}/C{i-1}-1,0)").number_format = "0.0%;-0.0%;-"
        for c in (4, 5):
            ws.cell(i, c).font = Font(name=F)
    vm_rango = (i0, i1)

    # ---------------- Categorías (participación con fórmula)
    ws = wb.create_sheet("Categorías")
    cat = r["categorias"][["categoria", "pedidos", "ingreso_brl"]].rename(
        columns={"categoria": "Categoría", "pedidos": "Pedidos", "ingreso_brl": "Ingreso (BRL)"})
    c0, c1 = _tabla(ws, cat, {2: "#,##0", 3: "#,##0"}, [32, 12, 16, 16, 16],
                    "Top 15 categorías por ingreso",
                    "Participación sobre el ingreso TOTAL del periodo (todas las categorías).")
    _encabezado(ws, 4, ["Categoría", "Pedidos", "Ingreso (BRL)", "% del ingreso total", "% acumulado"])
    for i in range(c0, c1 + 1):
        ws.cell(i, 4, f"=C{i}/Dashboard!$C$7").number_format = "0.0%"
        ws.cell(i, 5, f"=SUM($C${c0}:C{i})/Dashboard!$C$7").number_format = "0.0%"
        for c in (4, 5):
            ws.cell(i, c).font = Font(name=F)
    cat_rango = (c0, c1)

    # ---------------- Entrega vs calificación
    ws = wb.create_sheet("Entregas")
    ev = r["entrega_vs_calificacion"].copy()
    ev["rango_entrega"] = ev["rango_entrega"].str[3:]
    ev["pct_calif_1_o_2"] = ev["pct_calif_1_o_2"] / 100
    ev.columns = ["Entrega respecto a la fecha prometida", "Pedidos", "Calificación promedio (1-5)",
                  "% con calificación 1 o 2"]
    e0, e1 = _tabla(ws, ev, {2: "#,##0", 3: "0.00", 4: "0.0%"}, [36, 12, 18, 18],
                    "Puntualidad de la entrega vs. calificación del cliente",
                    "Excluye 189 pedidos con fechas inconsistentes y pedidos sin reseña.")
    ent_rango = (e0, e1)

    # ---------------- Estados
    ws = wb.create_sheet("Estados")
    es = r["estados"].copy()
    es.insert(1, "nombre", es["estado"].map(ESTADOS))
    es["pct_entregas_tardias"] = es["pct_entregas_tardias"] / 100
    es.columns = ["Estado", "Nombre", "Pedidos", "Ingreso (BRL)", "Días de entrega (promedio)",
                  "% entregas tardías", "Calificación promedio"]
    _tabla(ws, es, {3: "#,##0", 4: "#,##0", 5: "0.0", 6: "0.0%", 7: "0.00"},
           [9, 22, 11, 15, 16, 14, 14], "Desempeño por estado del cliente",
           "Ordenado por ingreso. Excluye pedidos con fechas inconsistentes.")

    # ---------------- Recompra
    ws = wb.create_sheet("Recompra")
    rc = r["recompra"][["pedidos_por_cliente", "clientes", "gasto_promedio_brl"]]
    rc.columns = ["Pedidos por cliente", "Clientes", "Gasto promedio (BRL)"]
    r0, r1 = _tabla(ws, rc, {2: "#,##0", 3: "#,##0.00"}, [20, 12, 18, 16],
                    "¿Cuántos clientes vuelven a comprar?",
                    "Clientes identificados con customer_unique_id (no customer_id, que cambia en cada pedido).")
    _encabezado(ws, 4, ["Pedidos por cliente", "Clientes", "Gasto promedio (BRL)", "% de clientes"])
    for i in range(r0, r1 + 1):
        ws.cell(i, 4, f"=B{i}/SUM($B${r0}:$B${r1})").number_format = "0.00%"
        ws.cell(i, 4).font = Font(name=F)
    rec_rango = (r0, r1)

    # ---------------- Calidad de datos y KPIs de SQL
    ws = wb.create_sheet("Calidad de datos")
    q = r["calidad"].rename(columns={"concepto": "Concepto", "registros": "Registros"})
    _tabla(ws, q, {2: "#,##0"}, [58, 14], "Problemas encontrados en los datos crudos",
           "Ninguna fila se borró en silencio: cada exclusión está marcada en 01_crear_base.sql.")
    k = r["kpis"].iloc[0]

    # ---------------- Dashboard
    d = dash
    d.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGHIJKLMNOP", [2, 4, 18, 3, 18, 3, 18, 3, 18, 3, 18, 3, 18, 3, 3, 3]):
        d.column_dimensions[col].width = w
    d["B2"] = "Dashboard de ventas · E-commerce Olist (Brasil)"
    d["B2"].font = Font(name=F, bold=True, size=18)
    d["B3"] = "Pedidos entregados de enero 2017 a agosto 2018 · montos en reales brasileños (BRL)"
    d["B3"].font = Font(name=F, size=10, color=TINTA2)

    vm_ref = f"'Ventas mensuales'!$C${vm_rango[0]}:$C${vm_rango[1]}"
    vp_ref = f"'Ventas mensuales'!$B${vm_rango[0]}:$B${vm_rango[1]}"
    tiles = [
        ("C", "Ingreso total", f"=SUM({vm_ref})", "#,##0"),
        ("E", "Pedidos", f"=SUM({vp_ref})", "#,##0"),
        ("G", "Ticket promedio", "=C7/E7", "#,##0.00"),
        ("I", "Entregas tardías", float(k.pct_entregas_tardias) / 100, "0.0%"),
        ("K", "Calificación promedio", float(k.calificacion_promedio), "0.00"),
        ("M", "Clientes que recompran", f"=1-Recompra!D{rec_rango[0]}", "0.0%"),
    ]
    caja = PatternFill("solid", start_color=GRIS)
    borde = Border(top=Side(style="thick", color=AZUL))
    for col, etiqueta, valor, fmt in tiles:
        e, v = d[f"{col}6"], d[f"{col}7"]
        e.value, v.value = etiqueta, valor
        e.font = Font(name=F, size=9, color=TINTA2)
        v.font = Font(name=F, size=18, bold=True)
        v.number_format = fmt
        for c in (e, v, d[f"{col}8"]):
            c.fill = caja
        e.border = borde
    d["I8"] = "Calculado en SQL"
    d["K8"] = "Calculado en SQL"
    for c in ("I8", "K8"):
        d[c].font = Font(name=F, size=7, italic=True, color=TINTA2)
    d.row_dimensions[7].height = 30

    # Gráficas nativas (se actualizan si cambian las tablas)
    lc = LineChart()
    lc.title, lc.height, lc.width = "Ingreso mensual (BRL)", 7.5, 13.5
    lc.y_axis.title, lc.y_axis.numFmt, lc.legend = None, "#,##0", None
    ws_vm = wb["Ventas mensuales"]
    lc.add_data(Reference(ws_vm, min_col=3, min_row=4, max_row=vm_rango[1]), titles_from_data=True)
    lc.set_categories(Reference(ws_vm, min_col=1, min_row=vm_rango[0], max_row=vm_rango[1]))
    lc.series[0].graphicalProperties.line.solidFill = AZUL
    lc.series[0].graphicalProperties.line.width = 28000
    lc.series[0].smooth = False
    d.add_chart(lc, "B10")

    bc = BarChart()
    bc.type, bc.title, bc.height, bc.width = "bar", "Top 10 categorías por ingreso (BRL)", 7.5, 16
    bc.legend = None
    ws_c = wb["Categorías"]
    bc.add_data(Reference(ws_c, min_col=3, min_row=4, max_row=cat_rango[0] + 9), titles_from_data=True)
    bc.set_categories(Reference(ws_c, min_col=1, min_row=cat_rango[0], max_row=cat_rango[0] + 9))
    bc.x_axis.scaling.orientation = "maxMin"
    bc.y_axis.numFmt = "#,##0"
    bc.series[0].graphicalProperties.solidFill = AZUL
    d.add_chart(bc, "I10")

    ec = BarChart()
    ec.title, ec.height, ec.width = "Calificación según puntualidad", 7.5, 13.5
    ec.legend = None
    ws_e = wb["Entregas"]
    ec.add_data(Reference(ws_e, min_col=3, min_row=4, max_row=ent_rango[1]), titles_from_data=True)
    ec.set_categories(Reference(ws_e, min_col=1, min_row=ent_rango[0], max_row=ent_rango[1]))
    ec.y_axis.scaling.min, ec.y_axis.scaling.max = 0, 5
    ec.series[0].graphicalProperties.solidFill = "EB6834"
    ec.dataLabels = DataLabelList(); ec.dataLabels.showVal = True
    ec.dataLabels.showSerName = False; ec.dataLabels.showCatName = False; ec.dataLabels.showLegendKey = False
    d.add_chart(ec, "B26")

    d["J27"] = "Hallazgos"
    d["J27"].font = Font(name=F, bold=True, size=12)
    hallazgos = [
        "1. Una entrega con más de 7 días de retraso baja la calificación de 4.3 a 1.7.",
        "2. El 97% de los clientes compra una sola vez: la recompra es la mayor oportunidad.",
        "3. Las 5 categorías principales generan el 40% del ingreso.",
        "4. Noviembre 2017 (Black Friday) fue el mes récord: +52% vs. octubre.",
        "5. São Paulo concentra el 38% del ingreso y entrega en 8.7 días en promedio.",
    ]
    for n, h in enumerate(hallazgos):
        c = d.cell(28 + n * 2, 10, h)
        c.font = Font(name=F, size=10)
    d["B42"] = FUENTE
    d["B42"].font = Font(name=F, size=8, italic=True, color=TINTA2)
    d["B43"] = "Análisis: Leonardo Chapa Gallardo · SQL (SQLite) + Python"
    d["B43"].font = Font(name=F, size=8, italic=True, color=TINTA2)

    d.page_setup.orientation = 'landscape'
    d.page_setup.fitToWidth = 1
    d.page_setup.fitToHeight = 1
    d.sheet_properties.pageSetUpPr.fitToPage = True
    for ws in wb.worksheets[1:]:
        ws["A3"] = FUENTE
        ws["A3"].font = Font(name=F, size=8, italic=True, color=TINTA2)
    wb.save(ruta)
